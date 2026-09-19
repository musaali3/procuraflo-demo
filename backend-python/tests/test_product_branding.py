import asyncio
import io
import zipfile
from pathlib import Path
from app.routes import settings


def test_registered_product_brand_preserves_company_identity(monkeypatch):
    monkeypatch.setattr(settings, 'fetch_one', lambda _: {'name': 'Independent Company', 'logo_url': '/uploads/logos/company.png'})
    result = settings.branding()
    assert result['application_name'] == 'Procuraflo'
    assert result['product_brand']['name'] == 'Procuraflo'
    assert 'procuraflo-logo.png' in result['product_brand']['logo_url']
    assert result['company_name'] == 'Independent Company'
    assert result['logo_url'] == '/uploads/logos/company.png'


def test_import_workbook_embeds_logo_and_preserves_import_columns():
    template_type = next(iter(settings.IMPORT_TEMPLATES))
    response = settings.download_import_template(template_type, {})
    async def collect():
        return b''.join([chunk async for chunk in response.body_iterator])
    content = asyncio.run(collect())
    with zipfile.ZipFile(io.BytesIO(content)) as book:
        media = [name for name in book.namelist() if name.startswith('xl/media/')]
        assert len(media) == 1
        source = Path(__file__).resolve().parents[2] / 'frontend-js/public/branding/procuraflo-logo.png'
        assert book.read(media[0]) == source.read_bytes()
        assert b'Procuraflo' in book.read('docProps/core.xml')
        assert b'roundRect' in book.read('xl/drawings/drawing1.xml')
    from openpyxl import load_workbook
    sheet = load_workbook(io.BytesIO(content)).active
    assert list(next(sheet.iter_rows(min_row=3, max_row=3, values_only=True))) == settings.IMPORT_TEMPLATES[template_type][0]
    assert sheet.freeze_panes == 'A4'
    assert sheet.print_title_rows == '$1:$3'
    assert not sheet.oddFooter.center.text
